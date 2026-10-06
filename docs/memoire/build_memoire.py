from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1]
OUT=ROOT/'output/documents';OUT.mkdir(parents=True,exist_ok=True)
text=(HERE/'content.md').read_text(encoding='utf-8')
TR=ROOT/'results/training/smartbug-403bb881-20260922T101754Z'
ev=json.loads((TR/'evaluation.json').read_text())
def pct(v):return f'{100*v:.2f}'.replace('.',',')
TABLES={
'vulnerabilities':('1.1','Familles de vulnérabilités considérées.', ['Famille','Principe du risque'],[
['Arithmetic','Erreurs liées aux opérations et limites numériques'],['Unchecked low calls','Résultat d’un appel de bas niveau non vérifié'],['Reentrancy','Réentrée pendant une interaction externe'],['Denial of service','Blocage ou indisponibilité d’une fonction'],['Front running','Exploitation de l’ordre des transactions'],['Time manipulation','Dépendance fragile à l’horodatage'],['Bad randomness','Aléa prévisible ou influençable']],[34,66]),
'metrics':('1.2','Définition des indicateurs issus de la confusion.',['Indicateur','Expression'],[
['Exactitude','(VP + VN) / (VP + VN + FP + FN)'],['Précision positive','VP / (VP + FP)'],['Rappel positif','VP / (VP + FN)'],['F1 positif','2 × précision × rappel / (précision + rappel)'],['FPR','FP / (FP + VN)']],[35,65]),
'literature':('1.3','Démarches des travaux connexes.',['Travail','Représentation / démarche'],[
['Tann et al. (2018)','Apprentissage séquentiel LSTM'],['Zhuang et al. (2020)','Réseau neuronal sur graphe'],['Zhang et al. (2022)','CBGRU, modèle hybride'],['Feng et al. (2020)','CodeBERT, encodeur préentraîné'],['Chen et Guestrin (2016)','XGBoost, arbres par boosting'],['Ferreira et al. (2020)','Framework SmartBugs']],[38,62]),
'partitions':('2.1','Composition du corpus final par partition.',['Partition','Total','Label 0','Label 1','Groupes'],[
['Entraînement','11 959','6 072','5 887','9 196'],['Validation','1 709','867','842','1 310'],['Calibration','1 709','868','841','1 235'],['Test','1 709','867','842','1 295'],['Source réservée','1 152','6','1 146','1 054'],['Total','18 238','8 680','9 558','14 090']],[32,17,17,17,17]),
'files':('2.2','Principaux fichiers et responsabilités.',['Fichier ou dossier','Rôle'],[
['src/build_dataset.py','Nettoyage, groupes et partitions'],['src/audit_dataset.py','Audit indépendant des représentations'],['src/preprocessing.py','Tokenisation et encodage communs'],['src/model.py','Construction du CNN–BiLSTM actif'],['src/train.py','Entraînement, validation, calibration et test'],['src/experiment.py','Lots et calcul des métriques'],['config/model.json','Hyperparamètres du modèle actif'],['src/predictor.py','Inférence commune au web et à la commande'],['src/run_model_benchmark.py','Comparaison CNN–BiLSTM / XGBoost / CodeBERT'],['config/benchmark_models.json','Réglages des concurrents et protocole']],[50,50]),
'hyperparameters':('3.1','Réglages du CNN–BiLSTM actif.',['Paramètre','Valeur'],[
['Graines initiales / retenue','42, 73, 101 / 73'],['Embedding / filtres','32 / 24'],['Noyau / unités BiLSTM','5 / 24 par sens'],['Dropout / dense','0,30 / 32'],['Optimiseur / taux','Adam / 0,001'],['Époques max. / patience','12 / 3'],['Lot de référence','32, réduit selon longueur'],['Threads CPU','4']],[45,55]),
'active_results':('3.2','Résultats du test pour le CNN–BiLSTM actif.',['Méthode','Exact. (%)','F1 macro (%)','Rappel (%)'],[],[37,21,21,21]),
'comparison':('3.3','État de la comparaison à trois modèles.',['Méthode','F1 macro (%)','Rappel (%)','FPR (%)'],[],[34,26,20,20]),
}
for name,k in [('CNN–BiLSTM brut','neural_raw'),('CNN–BiLSTM calibré','neural_calibrated')]:
    r=ev['test'][k]; TABLES['active_results'][3].append([name,pct(r['accuracy']),pct(r['f1_macro']),pct(r['recall_vulnerable'])])
r=ev['test']['neural_calibrated'];tn,fp,fn,tp=[n for row in r['confusion_matrix'] for n in row]
TABLES['comparison'][3].append(['CNN–BiLSTM actif',pct(r['f1_macro']),pct(r['recall_vulnerable']),pct(fp/(tn+fp))])
snapshot=json.loads((HERE/'comparison_snapshot.json').read_text(encoding='utf-8'))
source=ROOT/snapshot['source']
import hashlib
assert hashlib.sha256(source.read_bytes()).hexdigest()==snapshot['source_sha256'],'Changed XGBoost evaluation'
r=snapshot['metrics']
TABLES['comparison'][3].append(['XGBoost / TF-IDF',pct(r['f1_macro']),pct(r['recall_vulnerable']),pct(r['false_positive_rate'])])
TABLES['comparison'][3].append(['CodeBERT (candidat)','Calcul à compléter','Calcul à compléter','Calcul à compléter'])

ABBREVS=[('API','Application Programming Interface : interface de programmation'),('AUC','Area Under the Curve : aire sous la courbe'),('BiLSTM','Bidirectional Long Short-Term Memory : LSTM bidirectionnel'),('CBGRU','Nom du modèle hybride de Zhang et al. (2022)'),('CGT','Consolidated Ground Truth : vérité terrain consolidée'),('CNN','Convolutional Neural Network : réseau neuronal convolutif'),('CodeBERT','Encodeur préentraîné pour le code et le langage naturel'),('CPU','Central Processing Unit : processeur central'),('CSS','Cascading Style Sheets : feuilles de style en cascade'),('ECE','Expected Calibration Error : erreur de calibration estimée'),('EVM','Ethereum Virtual Machine : machine virtuelle Ethereum'),('F1','Moyenne harmonique de la précision et du rappel'),('FN','Faux négatif'),('FP','Faux positif'),('FPR','False Positive Rate : taux de faux positifs'),('HTML','HyperText Markup Language : langage de balisage hypertexte'),('IA','Intelligence artificielle'),('IC','Intervalle de confiance'),('JSON','JavaScript Object Notation : format de données structurées'),('LSTM','Long Short-Term Memory : mémoire récurrente à long et court terme'),('ReLU','Rectified Linear Unit : unité linéaire rectifiée'),('ROC','Receiver Operating Characteristic : courbe rappel / faux positifs'),('SHA-256','Secure Hash Algorithm, empreinte cryptographique de 256 bits'),('SQL','Structured Query Language : langage de requêtes structurées'),('TF-IDF','Term Frequency–Inverse Document Frequency : pondération lexicale'),('UML','Unified Modeling Language : langage de modélisation unifié'),('UTF-8','Unicode Transformation Format, encodage en unités de 8 bits'),('VN','Vrai négatif'),('VP','Vrai positif')]
BIB=[
('Livres',[
('GOODFELLOW, I., BENGIO, Y. et COURVILLE, A. (2016). Deep Learning. MIT Press.','https://www.deeplearningbook.org/','deeplearningbook.org'),]),
('Articles et communications',[
('FEIST, J., GRIECO, G. et GROCE, A. (2019). Slither: A Static Analysis Framework for Smart Contracts. WETSEB. DOI : 10.1109/WETSEB.2019.00008.','https://arxiv.org/abs/1908.09878','arxiv.org/abs/1908.09878'),
('FERREIRA, J. F., CRUZ, P., DURIEUX, T. et ABREU, R. (2020). SmartBugs: A Framework to Analyze Solidity Smart Contracts. Prépublication arXiv:2007.04771.','https://arxiv.org/abs/2007.04771','arxiv.org/abs/2007.04771'),
('GUO, C., PLEISS, G., SUN, Y. et WEINBERGER, K. Q. (2017). On Calibration of Modern Neural Networks. PMLR, 70, 1321–1330.','https://proceedings.mlr.press/v70/guo17a.html','proceedings.mlr.press/v70/guo17a'),
('KIM, Y. (2014). Convolutional Neural Networks for Sentence Classification. EMNLP, 1746–1751. DOI : 10.3115/v1/D14-1181.','https://aclanthology.org/D14-1181/','aclanthology.org/D14-1181'),
('TANN, W. J.-W., HAN, X. J., SEN GUPTA, S. et ONG, Y.-S. (2018). Towards Safer Smart Contracts: A Sequence Learning Approach to Detecting Security Threats. Prépublication arXiv:1811.06632, révisée en 2019.','https://arxiv.org/abs/1811.06632','arxiv.org/abs/1811.06632'),]),
('Articles et communications (suite)',[
('FENG, Z. et al. (2020). CodeBERT: A Pre-Trained Model for Programming and Natural Languages. Findings of EMNLP, 1536–1547. DOI : 10.18653/v1/2020.findings-emnlp.139.','https://aclanthology.org/2020.findings-emnlp.139/','aclanthology.org/2020.findings-emnlp.139'),
('CHEN, T. et GUESTRIN, C. (2016). XGBoost: A Scalable Tree Boosting System. KDD, 785–794. DOI : 10.1145/2939672.2939785.','https://arxiv.org/abs/1603.02754','arxiv.org/abs/1603.02754'),
('ZHANG, L. et al. (2022). CBGRU: A Detection Method of Smart Contract Vulnerability Based on a Hybrid Model. Sensors, 22(9), 3577. DOI : 10.3390/s22093577.','https://doi.org/10.3390/s22093577','doi.org/10.3390/s22093577'),
('ZHUANG, Y., LIU, Z., QIAN, P., LIU, Q., WANG, X. et HE, Q. (2020). Smart Contract Vulnerability Detection using Graph Neural Network. IJCAI, 3283–3290. DOI : 10.24963/ijcai.2020/454.','https://www.ijcai.org/Proceedings/2020/454','ijcai.org/Proceedings/2020/454'),]),
('Documentation et ressources du projet',[
('SALZER, G. et contributeurs (s. d.). Consolidated Ground Truth (CGT) for Weaknesses of Ethereum Smart Contracts. Dépôt de données.','https://github.com/gsalzer/cgt','github.com/gsalzer/cgt'),
('SOLIDITY (s. d.). Security Considerations. Documentation officielle.','https://docs.soliditylang.org/en/latest/security-considerations.html','docs.soliditylang.org — Security Considerations'),
('KOMPANI KIPANGU, E. (2026). SMART BUG : code, configurations, corpus préparé et résultats expérimentaux. Dépôt du mémoire, commit 2df2580.','https://github.com/Exauce-kompany/DETECTION-AUTOMATISEE-DES-VULNERABILITES-DE-SMARTS-CONTRACTS-SOLIDITY/tree/2df2580cf206bc46cb23864fda32d6db6c1c09c6','GitHub — dépôt SMART BUG, état étudié'),])]
ABBREVS += [('AdamW','Adam avec décroissance des poids découplée'),('BERT','Bidirectional Encoder Representations from Transformers'),('BPE','Byte Pair Encoding : encodage en sous-tokens'),('CLS','Classification token : marqueur de représentation de fenêtre'),('L2','Régularisation quadratique'),('PR-AUC','Precision–Recall Area Under the Curve : aire précision-rappel'),('XGBoost','eXtreme Gradient Boosting : boosting d’arbres'),('DOI','Digital Object Identifier : identifiant pérenne d’une publication'),('EMNLP','Empirical Methods in Natural Language Processing'),('IJCAI','International Joint Conference on Artificial Intelligence'),('MIT','Massachusetts Institute of Technology'),('PMLR','Proceedings of Machine Learning Research'),('s. d.','Sans date de publication précisée'),('WETSEB','Workshop on Emerging Trends in Software Engineering for Blockchain')]
ABBREVS += [('KDD','Knowledge Discovery and Data Mining : conférence sur l’analyse des données')]
ABBREVS.sort(key=lambda pair:pair[0].casefold())

doc=Document();WIDTH=398;FONT='Palatino Linotype'
for name,size,bold in [('Normal',12,False),('Title',20.7,True),('Heading 1',24.8,True),('Heading 2',16,True),('Heading 3',13,True),('Caption',10,False)]:
    s=doc.styles[name];s.font.name=FONT;s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    r=s.element.get_or_add_rPr();r.rFonts.set(qn('w:ascii'),FONT);r.rFonts.set(qn('w:hAnsi'),FONT)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme'):r.rFonts.attrib.pop(qn('w:'+key),None)
    pf=s.paragraph_format;pf.line_spacing=Pt(17.45);pf.space_after=Pt(11);pf.space_before=Pt(0);pf.widow_control=True
    for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    pf.first_line_indent=Pt(12 if name=='Normal' else 0);pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if name=='Normal' else WD_ALIGN_PARAGRAPH.LEFT
    if name.startswith('Heading'):pf.keep_with_next=True;pf.space_before=Pt(10);pf.space_after=Pt(14);pf.line_spacing=1.15
for name,size in [('Small',10),('Bibliography',10.5),('Contents',11),('Abbreviation',10.5),('Chapter Label',24.8)]:
    s=doc.styles.add_style(name,1);s.base_style=doc.styles['Normal'];s.font.size=Pt(size);s.paragraph_format.first_line_indent=Pt(0);s.paragraph_format.line_spacing=Pt(13 if name!='Chapter Label' else 30);s.paragraph_format.space_after=Pt(7 if name=='Bibliography' else 4)
doc.styles['Chapter Label'].font.bold=True
doc.styles['Small'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0);doc.styles['Caption'].paragraph_format.line_spacing=Pt(12);doc.styles['Caption'].paragraph_format.space_after=Pt(10)
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'fr-FR');doc.styles['Normal'].element.get_or_add_rPr().append(lang)
doc.core_properties.title='Détection automatisée des vulnérabilités des smart contracts Solidity : une approche basée sur l’apprentissage profond'
doc.core_properties.author='Exauce KOMPANI KIPANGU';doc.core_properties.subject='Mémoire de master — SMART BUG'
settings=doc.settings.element;hyph=OxmlElement('w:autoHyphenation');hyph.set(qn('w:val'),'true');settings.append(hyph)

def p(t='',style=None):return doc.add_paragraph(t,style)
def bookmark(par,name):
    st=OxmlElement('w:bookmarkStart');st.set(qn('w:id'),str(bookmark.n));st.set(qn('w:name'),name);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark.n));bookmark.n+=1;par._p.insert(0,st);par._p.append(end)
bookmark.n=1
def link(par,label,url=None,anchor=None):
    h=OxmlElement('w:hyperlink')
    if url:h.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
    else:h.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),FONT);fonts.set(qn('w:hAnsi'),FONT);pr.append(fonts);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
def set_page(label,header=''):
    sec=doc.sections[0] if label=='cover' else doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.page_width=Pt(595.84);sec.page_height=Pt(842.74);sec.left_margin=sec.right_margin=Pt(98.92);sec.top_margin=Pt(79);sec.bottom_margin=Pt(95);sec.header_distance=Pt(43);sec.footer_distance=Pt(91)
    sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
    hp=sec.header.paragraphs[0];hp.paragraph_format.first_line_indent=Pt(0);hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;hp.paragraph_format.space_after=Pt(4)
    if header:
        r=hp.add_run(header);r.font.size=Pt(10);r.italic=True
        bd=OxmlElement('w:pBdr');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:space'),'3');bd.append(b);hp._p.get_or_add_pPr().append(bd)
    fp=sec.footer.paragraphs[0];fp.paragraph_format.first_line_indent=Pt(0);fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.space_after=Pt(0)
    if label!='cover':
        roman=['i','ii','iii','iv','v','vi','vii','viii','ix'];num=roman.index(label)+1 if label in roman else int(label)
        pn=OxmlElement('w:pgNumType');pn.set(qn('w:start'),str(num));pn.set(qn('w:fmt'),'lowerRoman' if label in roman else 'decimal');sec._sectPr.append(pn)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');rp.append(sz);r.append(rp);tx=OxmlElement('w:t');tx.text=label;r.append(tx);fld.append(r);fp._p.append(fld)
    return sec
def major(t,chapter=None):
    if chapter:
        par=p('Chapitre '+chapter,'Chapter Label');par.paragraph_format.space_before=Pt(100);par.paragraph_format.space_after=Pt(36);par.paragraph_format.keep_with_next=True
        par=p(t,'Heading 1');par.paragraph_format.space_before=Pt(0);par.paragraph_format.space_after=Pt(40)
    else:
        par=p(t,'Heading 1');par.paragraph_format.space_before=Pt(70 if t=='Liste des abréviations' else 98);par.paragraph_format.space_after=Pt(30 if t=='Liste des abréviations' else 48)
    bookmark(par,'page_'+page_label)
def table(key):
    n,title,headers,rows,ratios=TABLES[key]
    cap=p('Tableau '+n+' – '+title,'Caption');cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=[round(WIDTH*20*r/100) for r in ratios];widths[-1]=round(WIDTH*20)-sum(widths[:-1]);pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:w'),str(sum(widths)));tw.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'90');ind.set(qn('w:type'),'dxa');pr.append(ind)
    margins=OxmlElement('w:tblCellMar')
    for k,v in [('top',65),('bottom',65),('left',90),('right',90)]:
        el=OxmlElement('w:'+k);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins);borders=OxmlElement('w:tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+k);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'888888');borders.append(el)
    pr.append(borders)
    for grid,w in zip(t._tbl.tblGrid.gridCol_lst,widths):grid.set(qn('w:w'),str(w))
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row();trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for cell,w,v in zip(row.cells,widths,values):
            cell.width=Pt(w/20);tcw=cell._tc.get_or_add_tcPr().find(qn('w:tcW'));tcw.set(qn('w:w'),str(w));tcw.set(qn('w:type'),'dxa');cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par=cell.paragraphs[0];par.paragraph_format.first_line_indent=Pt(0);par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=Pt(12);par.alignment=WD_ALIGN_PARAGRAPH.LEFT
            run=par.add_run(v);run.font.size=Pt(9.5);run.bold=i==0
    source=p('Source : '+('élaboration personnelle à partir des références citées.' if key in ['vulnerabilities','metrics','literature'] else 'configuration, données et résultats archivés de SMART BUG.'),'Small');source.paragraph_format.space_after=Pt(10)
    if key=='comparison':
        run=source.add_run('\nCNN et XGBoost : checkpoints sélectionnés sur validation. CodeBERT : candidat non entraîné.');run.font.size=Pt(9)
def figure(name,n,caption):
    par=p();par.paragraph_format.first_line_indent=Pt(0);par.paragraph_format.space_after=Pt(4);par.paragraph_format.keep_with_next=True;par.paragraph_format.line_spacing=1.0
    par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(HERE/'figures'/f'{name}.png'),width=Pt(275 if name=='confusion' else WIDTH))
    c=p('Figure '+n+' – '+caption,'Caption');c.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=c.add_run('\nSource : '+('élaboration personnelle d’après le système implémenté.' if n.startswith('2.') else 'élaboration personnelle.' if n.startswith('1.') else 'protocole comparatif ; évaluation CodeBERT à compléter.' if name in ['comparison','ablation','latency'] else 'résultats expérimentaux archivés de SMART BUG.'));r.font.size=Pt(9)

pages=re.split(r'^@PAGE ',text,flags=re.M)[1:]
figlist=[];tablelist=[];toc=[]
for block in pages:
    label,body=block.split('\n',1)
    for line in body.splitlines():
        if line.startswith('@FIG '):name,n,title=line[5:].split('|');figlist.append((n+' '+title,label))
        elif line.startswith('@TABLE '):n,title,*_=TABLES[line[7:]];tablelist.append((n+' '+title,label))
        elif line.startswith('@MAJOR '):
            title=line[7:]
            if title!='Table des matières':toc.append((title,label,0))
        elif line.startswith('@CHAPTER '):n,title=line[9:].split('|');toc.append((n+' '+title,label,0))
        elif line.startswith('## '):toc.append((line[3:],label,1))
        elif line.startswith('### '):toc.append((line[4:],label,2))
toc.insert(3,('Table des matières','v',0))
def entry(title,number,level=0):
    par=p(style='Contents');par.paragraph_format.left_indent=Pt(level*10);par.paragraph_format.line_spacing=Pt(14);par.paragraph_format.space_after=Pt(5)
    par.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    link(par,title,anchor='page_'+number);par.add_run('\t'+number)
def bib(groups):
    for head,entries in groups:
        par=p(head,'Heading 3');par.paragraph_format.space_before=Pt(8);par.paragraph_format.space_after=Pt(8)
        for description,url,label in entries:
            par=p(description+' ','Bibliography');par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;link(par,label,url=url)

for block in pages:
    page_label,body=block.split('\n',1);header=re.search(r'^@HEADER (.*)$',body,re.M);sec=set_page(page_label,header.group(1) if header else '')
    if not re.search(r'^@(MAJOR|CHAPTER)',body,re.M) and page_label!='cover':bookmark(p('','Small'),'page_'+page_label)
    for para in re.split(r'\n\s*\n',body.strip()):
        # Directives may be adjacent; prose paragraphs are single blocks.
        if para.startswith('@') or para.startswith('#'):
            for line in para.splitlines():
                if not line or line.startswith('@HEADER '):continue
                if line=='@COVER':
                    par=p('Détection automatisée des vulnérabilités\ndes smart contracts Solidity :\nune approche basée sur\nl’apprentissage profond','Title');par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.space_before=Pt(68);par.paragraph_format.line_spacing=Pt(28);par.paragraph_format.space_after=Pt(43)
                    par=p('par');par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.first_line_indent=Pt(0);par.paragraph_format.space_after=Pt(30)
                    par=p('Exauce KOMPANI KIPANGU');par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.first_line_indent=Pt(0);par.runs[0].font.size=Pt(17.2)
                    par=p('Octobre 2026');par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.first_line_indent=Pt(0);par.paragraph_format.space_before=Pt(335)
                elif line.startswith('@MAJOR '):major(line[7:])
                elif line.startswith('@CHAPTER '):n,t=line[9:].split('|');major(t,n)
                elif line.startswith('@DEDICATION '):
                    par=p(line[12:]);par.paragraph_format.space_before=Pt(405)
                elif line.startswith('@FIG '):figure(*line[5:].split('|'))
                elif line.startswith('@TABLE '):table(line[7:])
                elif line.startswith('@TOC '):
                    entries=toc[:21] if line.endswith('1') else toc[21:]
                    for t,n,l in entries:entry(t,n,l)
                elif line=='@FIGLIST':
                    for t,n in figlist:entry(t,n)
                elif line=='@TABLELIST':
                    for t,n in tablelist:entry(t,n)
                elif line=='@ABBREVS':
                    for short,long in ABBREVS:
                        par=p(style='Abbreviation');par.paragraph_format.line_spacing=Pt(11.3);par.paragraph_format.space_after=Pt(.4);par.add_run(short+' : ').bold=True;par.add_run(long)
                elif line.startswith('@BIB '):
                    bib(BIB[:2] if line.endswith('1') else BIB[2:])
                    if line.endswith('2'):
                        par=p('Références du protocole comparatif mises à jour le 5 octobre 2026. Les graphiques et tableaux de résultats proviennent des artefacts locaux cités au chapitre 2.','Small')
                elif line.startswith('### '):p(line[4:],'Heading 3')
                elif line.startswith('## '):p(line[3:],'Heading 2')
        else:p(para.replace('\n',' '))

path=OUT/'Memoire_SMART_BUG_Exauce_Kompani.docx';doc.save(path)
(HERE/'document_index.json').write_text(json.dumps({'pages':len(pages),'figures':figlist,'tables':tablelist,'toc':toc,'words':len(text.split()),'bibliography_entries':sum(len(entries) for _,entries in BIB),'abbreviations':len(ABBREVS),'comparison_status':'cnn_and_xgboost_evaluated_codebert_incomplete'},ensure_ascii=False,indent=2),encoding='utf-8')
print(path);print('Pages planned:',len(pages),'Figures:',len(figlist),'Tables:',len(tablelist),'Words:',len(text.split()))
